# -*- coding: utf-8 -*-
"""
生成《多智能体（LLM-MAS）调研汇报》完整版 PPT。

结构：封面 + 领域坐标(Yan 2025) + 目录 + 8 篇论文 × 4 页 + 致谢 = 36 页
每篇论文固定四页：题名页 → 动机与难点 → 模型 → 实验结果
顺序：CAMEL → MetaGPT → AutoGen → ChatDev → MacNet → AgentVerse
      → Multiagent Debate → GeoLLM-Squad(遥感)

用法： python3 build_ppt.py
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
FIG = os.path.join(ROOT, "figures")
OUT = os.path.join(HERE, "多智能体调研汇报.pptx")

# ---------- 设计系统 ----------
NAVY = RGBColor(0x0F, 0x2A, 0x4A)
BLUE = RGBColor(0x1F, 0x6F, 0xB2)
TEAL = RGBColor(0x0E, 0x8F, 0x82)
AMBER = RGBColor(0xC8, 0x6A, 0x00)
CRIMSON = RGBColor(0xB3, 0x26, 0x1E)
INK = RGBColor(0x22, 0x27, 0x2E)
GREY = RGBColor(0x5B, 0x64, 0x70)
LIGHT = RGBColor(0xF2, 0xF5, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BORDER = RGBColor(0xDD, 0xE3, 0xEA)
WARM = RGBColor(0xFD, 0xF3, 0xE6)
# Use a CJK font available on macOS so Chinese text survives PDF export.
FONT = "STHeiti"

SW, SH = Inches(13.333), Inches(7.5)
M = Inches(0.62)
TOTAL = 36


def _p(layer, x, y, w, h):
    tb = layer.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def _run(p, text, size=14, bold=False, color=INK, italic=False):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = FONT
    return r


def rect(layer, x, y, w, h, fill=LIGHT, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    s = layer.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    s.adjustments[0] = 0.06
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    return s


def new_slide(prs, section="", page=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    if section:
        tf = _p(s, M, Inches(0.26), SW - 2 * M - Inches(1.4), Inches(0.3))
        _run(tf.paragraphs[0], section, size=11, bold=True, color=GREY)
    if page:
        tf = _p(s, SW - M - Inches(1.1), SH - Inches(0.46), Inches(1.1), Inches(0.3))
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        _run(p, f"{page:02d} / {TOTAL}", size=10.5, color=GREY)
    return s


def header(s, title, subtitle=None, accent=BLUE):
    rect(s, M, Inches(0.62), Inches(0.09), Inches(0.46), fill=accent)
    tf = _p(s, M + Inches(0.24), Inches(0.55), SW - 2 * M - Inches(0.3), Inches(0.55))
    _run(tf.paragraphs[0], title, size=25, bold=True, color=NAVY)
    if subtitle:
        tf = _p(s, M + Inches(0.24), Inches(1.08), SW - 2 * M - Inches(0.3), Inches(0.34))
        _run(tf.paragraphs[0], subtitle, size=13, color=GREY)
    return Inches(1.52)


def bullets(s, x, y, w, h, items, size=14, gap=8, bullet="▪ "):
    tf = _p(s, x, y, w, h)
    first = True
    for it in items:
        level, color, text = 0, None, it
        if isinstance(it, tuple):
            text = it[0]
            if len(it) > 1:
                level = it[1] or 0
            if len(it) > 2:
                color = it[2]
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(gap)
        p.line_spacing = 1.18
        mark = "" if text.endswith("：") else ("– " if level else bullet)
        _run(p, mark + text, size=size - level, bold=text.endswith("："),
             color=color or (INK if level == 0 else GREY))
    return tf


def figure(s, name, x, y, w, h, caption=None):
    path = os.path.join(FIG, name)
    ph = h if caption is None else h - Inches(0.28)
    if os.path.exists(path):
        iw, ih = Image.open(path).size
        scale = min(w / iw, ph / ih)
        pic = s.shapes.add_picture(
            path, x + (w - Emu(int(iw * scale))) // 2,
            y + (ph - Emu(int(ih * scale))) // 2,
            width=Emu(int(iw * scale)), height=Emu(int(ih * scale)))
        pic.line.color.rgb = BORDER
        pic.line.width = Pt(0.75)
    if caption:
        tf = _p(s, x, y + h - Inches(0.26), w, Inches(0.26))
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        _run(p, caption, size=9.5, color=GREY)


# ==================== 页面构建 ====================
def p_cover(prs, page):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, SW, SH, fill=NAVY)
    rect(s, 0, Inches(4.58), SW, Inches(0.055), fill=TEAL)
    tf = _p(s, M, Inches(2.02), SW - 2 * M, Inches(1.0))
    _run(tf.paragraphs[0], "多智能体（LLM-MAS）调研汇报", size=40, bold=True, color=WHITE)
    tf = _p(s, M, Inches(3.12), SW - 2 * M, Inches(1.2))
    _run(tf.paragraphs[0], "按协作机制选代表性论文：每篇讲清 动机/难点 → 模型 → 实验结果",
         size=18, color=RGBColor(0xBF, 0xD9, 0xEE))
    tf = _p(s, M, Inches(5.05), SW - 2 * M, Inches(0.9))
    _run(tf.paragraphs[0], "汇报人：____________　　2026 年 9 月",
         size=14, color=RGBColor(0x9F, 0xB8, 0xCC))
    return s


def p_yan(prs, page):
    s = new_slide(prs, "领域坐标", page)
    y = header(s, "领域坐标：LLM-MAS 的通信分类体系",
               "Yan et al. 2025 · Beyond Self-Talk · arXiv:2502.14321 · 后面 8 篇论文都能被这套框架定位",
               TEAL)
    figure(s, "yan_taxonomy.png", M + Inches(6.55), y - Inches(0.05),
           SW - M - (M + Inches(6.55)), Inches(4.75), "Yan 2025 分类树（p.3）")
    bullets(s, M, y, Inches(6.2), Inches(4.8), [
        "问题：LLM 本不是为相互通信设计的。现有综述多按应用或 Agent 能力分类，通信只被当成实现细节。",
        "本文把通信提为主线，给出两级分类：系统级 + 系统内，共 7 个维度。",
        "系统级 · 架构：谁和谁说话 → 扁平 / 层级 / 团队 / 社会 / 混合。",
        "系统级 · 目标：为什么说话 → 合作（含辩论式）/ 竞争 / 混合。",
        "系统级 · 协议：消息怎么走 → MCP（连工具）/ A2A（找 Agent）/ ANP（开放网络）。",
        "系统内 · 策略：什么时候说 → 逐个 / 同时 / 同时+摘要器。",
        "系统内 · 范式与内容：怎么表示、说什么 → 消息传递 / 言语行为 / 黑板；显式 / 隐式。",
        "价值：它不只罗列分类，每个维度都给出代价（扁平难扩展、层级易过载、逐个会累积错误）。",
    ], size=12.5, gap=7)
    return s


def p_toc(prs, page, papers):
    s = new_slide(prs, "目录", page)
    y = header(s, "按协作机制逐篇精读", "每篇固定四页：题名页 → 动机与难点 → 模型 → 实验结果", BLUE)
    colw = (SW - 2 * M - Inches(0.3)) / 2
    for i, p in enumerate(papers):
        col, row = i // 4, i % 4
        x = M + col * (colw + Inches(0.3))
        yy = y + (Inches(1.06) + Inches(0.14)) * row
        rect(s, x, yy, colw, Inches(1.06), fill=LIGHT)
        rect(s, x, yy, Inches(0.07), Inches(1.06), fill=p["accent"])
        tf = _p(s, x + Inches(0.26), yy + Inches(0.12), colw - Inches(1.3), Inches(0.85))
        _run(tf.paragraphs[0], p["name"], size=17, bold=True, color=NAVY)
        p2 = tf.add_paragraph()
        p2.space_before = Pt(2)
        _run(p2, p["toc_sub"], size=12, color=GREY)
        tf = _p(s, x + colw - Inches(1.25), yy + Inches(0.3), Inches(1.0), Inches(0.4))
        pp = tf.paragraphs[0]
        pp.alignment = PP_ALIGN.RIGHT
        _run(pp, f"P{p['start']:02d}", size=13, bold=True, color=p["accent"])
    tf = _p(s, M, SH - Inches(0.98), SW - 2 * M, Inches(0.4))
    _run(tf.paragraphs[0], "共 8 篇论文 · 协作机制 7 篇（重点）+ 遥感落地 1 篇",
         size=12, color=GREY, italic=True)
    return s


def p_cover_paper(prs, page, p):
    s = new_slide(prs, p["section"], page)
    rect(s, M, Inches(1.55), Inches(0.09), Inches(0.5), fill=p["accent"])
    tf = _p(s, M + Inches(0.26), Inches(1.48), SW - 2 * M - Inches(0.4), Inches(0.6))
    _run(tf.paragraphs[0], p["name"], size=30, bold=True, color=NAVY)
    tf = _p(s, M, Inches(2.30), SW - 2 * M, Inches(1.4))
    for k, seg in enumerate(p["en_title"].split("\n")):
        pp = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        pp.line_spacing = 1.14
        _run(pp, seg, size=22, bold=True, color=p["accent"])
    tf = _p(s, M, Inches(3.85), SW - 2 * M, Inches(0.5))
    _run(tf.paragraphs[0], p["authors"], size=13, color=INK)
    tf = _p(s, M, Inches(4.24), SW - 2 * M, Inches(0.4))
    _run(tf.paragraphs[0], p["venue"], size=13, color=GREY, italic=True)
    rect(s, M, Inches(5.0), SW - 2 * M, Inches(1.0), fill=LIGHT)
    rect(s, M, Inches(5.0), Inches(0.075), Inches(1.0), fill=p["accent"])
    tf = _p(s, M + Inches(0.3), Inches(5.2), SW - 2 * M - Inches(0.6), Inches(0.65))
    _run(tf.paragraphs[0], p["one_liner"], size=16, bold=True, color=NAVY)
    return s


def p_motivation(prs, page, p):
    s = new_slide(prs, p["section"], page)
    accent = p["accent"]
    y = header(s, p["name"] + "｜动机与难点", p["motivation"]["headline"], accent)
    rect(s, M, y, SW - 2 * M, Inches(0.92), fill=LIGHT)
    rect(s, M, y, Inches(0.075), Inches(0.92), fill=accent)
    tf = _p(s, M + Inches(0.3), y + Inches(0.14), SW - 2 * M - Inches(0.6), Inches(0.68))
    pp = tf.paragraphs[0]
    _run(pp, "研究目的：", size=15, bold=True, color=accent)
    _run(pp, p["motivation"]["purpose"], size=15, color=NAVY)
    bullets(s, M, y + Inches(1.18), Inches(7.05), Inches(3.8),
            p["motivation"]["points"], size=13.5, gap=10)
    rect(s, M + Inches(7.4), y + Inches(1.18), Inches(4.65), Inches(3.8), fill=WARM)
    tf = _p(s, M + Inches(7.7), y + Inches(1.42), Inches(4.05), Inches(3.3))
    _run(tf.paragraphs[0], "难点（技术上的困难）", size=14.5, bold=True, color=AMBER)
    for c in p["motivation"]["challenges"]:
        pp = tf.add_paragraph()
        pp.space_before = Pt(10)
        pp.line_spacing = 1.14
        _run(pp, "· " + c, size=12.5, color=INK)
    return s


def p_model(prs, page, p):
    s = new_slide(prs, p["section"], page)
    accent = p["accent"]
    y = header(s, p["name"] + "｜模型", p["model"]["headline"], accent)
    figs = p["model"].get("figs") or []
    spec = p["model"]["bullets"]

    def ratio(f):
        iw, ih = Image.open(os.path.join(FIG, f)).size
        return iw / ih

    if figs and ratio(figs[0][0]) > 2.4:
        # 宽扁图：文字分两栏在上，图片占满整页宽度在下方
        half = (SW - 2 * M - Inches(0.4)) / 2
        cut = (len(spec) + 1) // 2
        bullets(s, M, y, half, Inches(2.0), spec[:cut], size=11.5, gap=7)
        bullets(s, M + half + Inches(0.4), y, half, Inches(2.0), spec[cut:], size=11.5, gap=7)
        top = y + Inches(2.18)
        f, c = figs[0]
        figure(s, f, M, top, SW - 2 * M, SH - Inches(0.72) - top, c)
    else:
        bullets(s, M, y, Inches(5.4), Inches(4.85), spec, size=12, gap=9)
        x2 = M + Inches(5.78)
        w2 = SW - M - x2
        top = y - Inches(0.05)
        avail = SH - Inches(0.78) - top
        if figs:
            f, c = figs[0]
            figure(s, f, x2, top, w2, avail, c)
        else:
            rect(s, x2, top, w2, Inches(2.4), fill=LIGHT)
            tf = _p(s, x2 + Inches(0.24), top + Inches(0.2), w2 - Inches(0.48), Inches(2.0))
            _run(tf.paragraphs[0], p["model"]["side_title"], size=13.5, bold=True, color=accent)
            for it in p["model"]["side_items"]:
                pp = tf.add_paragraph()
                pp.space_before = Pt(8)
                pp.line_spacing = 1.12
                _run(pp, it, size=12, color=INK)
    return s


def p_results_full(prs, page, p):
    """结果页 A：表格/图表占满宽度（放最大）。"""
    s = new_slide(prs, p["section"], page)
    y = header(s, p["name"] + "｜实验结果", p["results"]["headline"], p["accent"])
    top = y - Inches(0.05)
    avail = SH - Inches(0.72) - top
    items = p["results"]["items"]
    n = len(items)
    gap = Inches(0.3)
    cap_h = Inches(0.28)
    full_w = SW - 2 * M
    nat = []
    for f, _ in items:
        iw, ih = Image.open(os.path.join(FIG, f)).size
        nat.append(full_w / (iw / ih))
    budget = avail - gap * (n - 1) - cap_h * n
    scale = min(1.0, budget / sum(nat))
    yy = top
    for i, (f, c) in enumerate(items):
        h = nat[i] * scale + cap_h
        figure(s, f, M, yy, full_w, h, c)
        yy += h + gap
    return s


def p_results_split(prs, page, p):
    """结果页 B：主实验与消融并列两栏（无底色框），下方放一张图表。"""
    s = new_slide(prs, p["section"], page)
    accent = p["accent"]
    y = header(s, p["name"] + "｜实验结果", p["results"]["headline"], accent)
    half = (SW - 2 * M - Inches(0.4)) / 2
    for i, (title, items, color) in enumerate([
            ("主实验：与 SOTA 对比", p["results"]["main"], accent),
            ("动机实验（消融）", p["results"]["abl"], AMBER)]):
        x = M + i * (half + Inches(0.4))
        rect(s, x, y + Inches(0.04), Inches(0.06), Inches(0.3), fill=color)
        tf = _p(s, x + Inches(0.16), y, half - Inches(0.16), Inches(0.32))
        _run(tf.paragraphs[0], title, size=13.5, bold=True, color=color)
        bullets(s, x, y + Inches(0.42), half, Inches(1.85), items, size=11.5, gap=6)
    top = y + Inches(2.32)
    f, c = p["results"]["items"][0]
    figure(s, f, M, top, SW - 2 * M, SH - Inches(0.72) - top, c)
    return s


def p_thanks(prs, page):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, SW, SH, fill=NAVY)
    rect(s, 0, Inches(3.42), SW, Inches(0.06), fill=TEAL)
    tf = _p(s, M, Inches(2.55), SW - 2 * M, Inches(1.0))
    _run(tf.paragraphs[0], "感谢老师的指导", size=36, bold=True, color=WHITE)
    tf = _p(s, M, Inches(3.75), SW - 2 * M, Inches(0.8))
    _run(tf.paragraphs[0], "请老师指正", size=15, color=RGBColor(0xBF, 0xD9, 0xEE))
    return s


# ==================== 论文内容 ====================
PAPERS = [
    dict(
        name="CAMEL", accent=TEAL, section="协作机制",
        toc_sub="协作机制 · 角色扮演",
        en_title="Communicative Agents for “Mind” Exploration of\nLarge Language Model Society",
        venue="NeurIPS 2023 · arXiv:2303.17760",
        authors="Guohao Li, Hasan Abed Al Kader Hammoud, Hani Itani, Dmitrii Khizbullin, Bernard Ghanem　（KAUST）",
        one_liner="两个 Agent 靠角色设定自主协作：结构最简单、成本最低的一条路线。",
        motivation=dict(
            headline="两个 Agent 靠角色设定自主协作，不需要人一直引导",
            purpose="构造可扩展的 Agent 协作与数据生成机制，用于研究「Agent 社会」的行为。",
            points=[
                "复杂任务目前依赖人工一步步引导对话，任务越复杂，人机往返轮次越多，时间与人力成本随规模上升。",
                "人工介入还意味着无法批量产出协作数据，多智能体的行为难以被系统研究。",
                "论文要的不是更好的提示词，而是一套让两个 Agent 自主把任务推进到底的机制。",
            ],
            challenges=[
                "两个 LLM 之间没有天然的角色边界，容易发生角色反转（Assistant 反过来下指令）。",
                "对话缺少显式的停止条件，容易陷入无效循环或提前终止。",
                "生成的对话质量无法保证，会出现只复述指令、敷衍回复的情况。",
                "协作效果完全依赖底层模型的指令遵循能力，模型弱则机制失效。",
            ]),
        model=dict(
            headline="角色扮演 + 初始提示：人类只在最左边出现一次",
            bullets=[
                "角色分配：指定两个角色（如 Python 程序员 / 股票交易员），一个扮 AI User，一个扮 AI Assistant。",
                "Task Specifier：把人类一句模糊想法扩展成清晰、具体的任务提示词。",
                "AI User：提需求、给指令、推动进度——扮演人类。",
                "AI Assistant：执行、产出方案——扮演专家。",
                "指令跟随式对话：User 给指令 Iₜ → Assistant 给方案 Sₜ，多轮往返。",
                "终止标记：约定 CAMEL_TASK_DONE，出现即结束。",
                "关键点：人类只负责设定想法和角色，后面全靠两个 Agent 自己来回——这就是它省人力的地方。",
            ],
            figs=[("camel_framework.png", "Figure 1 CAMEL 角色扮演流程（NeurIPS 2023, p.4）")]),
        results=dict(
            headline="同一任务下双 Agent 协作显著优于单模型；生成的数据还能反哺小模型",
            items=[("tbl_camel_eval.png",
                    "Table 1 主实验：CAMEL 双 Agent vs gpt-3.5-turbo 单次回答（p.9）"),
                   ("tbl_camel_humaneval.png",
                    "Table 3 下游价值：用生成数据微调的 CAMEL-7B 在 HumanEval(+) 上的 pass@k（p.10）")],
            full=True)),

    dict(
        name="MetaGPT", accent=TEAL, section="协作机制",
        toc_sub="协作机制 · SOP 流水线",
        en_title="Meta Programming for A Multi-Agent\nCollaborative Framework",
        venue="ICLR 2024 · arXiv:2308.00352",
        authors="Sirui Hong, Mingchen Zhuge, Jiaqi Chen, Xiawu Zheng, Yuheng Cheng 等　（DeepWisdom / KAUST）",
        one_liner="把人类团队的 SOP 编码进多智能体：交文档、跑测试，最结构化的一条路线。",
        motivation=dict(
            headline="自由聊天会级联幻觉，需要把流程和中间产物写死",
            purpose="用结构化中间产物替代自由文本，用可执行反馈替代人工检查，把协作变成可复现的流程。",
            points=[
                "多个 Agent 自由聊天会出现信息歧义，一个 Agent 的错误成为下一个的输入，沿对话链放大（级联幻觉）。",
                "自然语言信息密度低，Agent 之间没有明确的「接口」，需求理解容易不一致。",
                "代码「看起来合理」不等于「能运行」，而流程中缺少自动验证环节。",
            ],
            challenges=[
                "如何把人类团队的 SOP 编码成 Agent 可执行的 Prompt、Role 与 Message。",
                "如何约束依赖关系：谁先工作、下游在什么条件下才允许启动。",
                "如何定义结构化中间产物，使它能被程序校验，而不只是被人阅读。",
                "角色越多质量越高但费用上升，需要在质量与开销之间做取舍。",
            ]),
        model=dict(
            headline="四个机制：角色分工 · SOP · 结构化通信 · 可执行反馈",
            bullets=[
                "Role 角色分工：产品经理 → 架构师 → 项目经理 → 工程师 → QA，每个角色有 profile / goal / constraints / skills / actions / memory。",
                "SOP 标准流程：规定谁先工作、依赖谁、输出必须含什么字段、何时可启动（架构师不能在收到 PRD 前开始设计）。",
                "结构化通信：Agent 之间传 PRD、系统设计、接口定义、任务列表、测试报告等中间交付物，而不是自由文本。",
                "Publish-Subscribe 共享消息池：Agent 发布结构化消息，其余按订阅关系取用，把连接数从 O(n²) 降到 O(n)。",
                "可执行反馈：生成代码 → 运行 / 单测 → 读错误 → 改代码 → 再运行，最多重试 3 次。",
            ],
            figs=[("meta_sop.png", "Figure 1 人类团队 SOP 与 MetaGPT 角色/产物的对应（ICLR 2024, p.2）")]),
        results=dict(
            headline="在代码生成与真实开发任务上同时优于 ChatDev 与单模型",
            items=[("tbl_metagpt_softwaredev.png",
                    "Table 1 SoftwareDev 统计：可执行性 3.75 / 4，人工修复成本 0.83（p.8）"),
                   ("tbl_metagpt_roles.png",
                    "Table 3 角色消融：逐步加入角色，可执行性 1.0 → 4.0（p.9）")],
            full=True)),

    dict(
        name="AutoGen", accent=TEAL, section="协作机制",
        toc_sub="协作机制 · 可编程对话",
        en_title="Enabling Next-Gen LLM Applications via\nMulti-Agent Conversation",
        venue="arXiv:2308.08155 · Microsoft Research",
        authors="Qingyun Wu, Gagan Bansal, Jieyu Zhang, Yiran Wu, Beibin Li 等",
        one_liner="把「Agent 之间怎么对话」变成可编程接口：协作拓扑成为设计变量。",
        motivation=dict(
            headline="把协作拓扑变成可编程的接口，而不是写死的流程",
            purpose="提供通用框架，让多个可对话 Agent 之间的交互成为一等公民，开发者能像写程序一样编排协作。",
            points=[
                "真实任务需要「模型 + 工具 + 人」混合参与，而写死的流程难以复用。",
                "让 Agent 完全自由聊天又不可控，无法保证任务被推进。",
                "开发者需要能自由定义对话模式，而不是接受一种固定的拓扑。",
            ],
            challenges=[
                "如何让 LLM、工具、人类在框架里成为同一种可互换的实体。",
                "如何在自然语言回复与程序化控制（执行代码、判断终止）之间自由切换。",
                "如何设计去中心化机制：没有中央调度器时，由谁推动下一步。",
                "如何让人类随时可介入，同时不破坏自动流程。",
            ]),
        model=dict(
            headline="Conversable Agent + 自动回复 = 可编程的对话",
            bullets=[
                "Conversable Agent：统一 generate_reply 接口，底层可以是 LLM、任意 Python 代码或人类代理。",
                "Auto-reply 自动回复：Agent 收到消息后自己决定下一步——继续回复、调用工具、或交出控制权。",
                "没有中央调度器，协作从「消息 + 自动回复」中涌现（去中心化、模块化）。",
                "Conversation Programming：开发者注册回复函数、设置终止条件来定义协作流程。",
                "支持的对话模式：双 Agent 对话、群聊、层级聊天、联合聊天。",
                "一句话：别人在设计 Agent 的智能，AutoGen 在设计 Agent 之间的对话。",
            ],
            figs=[("autogen_overview.png", "Figure 1 左：Agent 定制；中：灵活对话模式；右：完整对话实例（p.1）")]),
        results=dict(
            headline="六个应用都跑通；ALFWorld 加入 grounding agent 后平均 +15%",
            items=[("autogen_fig4_top.png",
                    "Figure 4(a)(b) 数学问题求解（左）与检索增强问答（右）（p.7）"),
                   ("autogen_fig4_bottom.png",
                    "Figure 4(c)(d)：左 ALFWorld；右 OptiGuide —— Multi = 多智能体，Single = 单智能体（p.7）")],
            full=True)),

    dict(
        name="ChatDev", accent=TEAL, section="协作机制",
        toc_sub="协作机制 · 聊天链",
        en_title="Communicative Agents for\nSoftware Development",
        venue="ACL 2024 · arXiv:2307.07924",
        authors="Chen Qian, Wei Liu, Hongzhang Liu, Nuo Chen, Yufan Dang 等　（清华大学等）",
        one_liner="用聊天链把开发流程拆成原子对话：语言本身就是协作协议。",
        motivation=dict(
            headline="用对话驱动整条开发流水线，规则与防幻觉都写进对话里",
            purpose="用同一种语言贯穿设计、编码与测试，让多智能体靠对话产出可运行的软件。",
            points=[
                "已有做法分别优化设计、编码、测试各阶段，导致技术不一致、流程碎片化。",
                "对话若没有流程约束，容易跳步或漏掉测试环节。",
                "Agent 在信息不足时会硬答，凭空编造外部依赖与接口（幻觉）。",
            ],
            challenges=[
                "如何把开发流程拆成可独立执行的原子子任务，并为每个子任务定义终止条件。",
                "如何在不打断流程的前提下，让 Agent 主动承认「信息不足」并追问。",
                "如何统一两种表达方式：设计阶段用自然语言，调试阶段用编程语言。",
                "如何保证多轮迭代收敛，而不是反复修改却不改进。",
            ]),
        model=dict(
            headline="Chat Chain：3 阶段 5 子任务 + 交际式去幻觉",
            bullets=[
                "Chat Chain：把开发拆成原子化子任务，每个子任务对应一次双 Agent 会话。",
                "3 阶段 5 子任务：Design → Coding → Testing；角色为 CEO、CTO、程序员、评审员、测试员。",
                "每个子任务内部是角色扮演式双 Agent 对话（Instructor 提要求 / Assistant 执行），继承自 CAMEL。",
                "终止条件：代码修改两轮不变，或超过 10 轮通信，即结束该子任务。",
                "产出以 <SOLUTION> 标记，便于程序自动抽取结果。",
                "交际式去幻觉：允许 Assistant 先反过来索要更具体信息再正式回答——「不确定就先问，不要先编」。",
            ],
            figs=[("chatdev_chain.png", "Figure 2 Chat Chain：阶段、子任务、角色与对话流（ACL 2024, p.3）")]),
        results=dict(
            headline="四项指标全面领先；消融显示测试环节与去幻觉机制是主要来源",
            items=[("tbl_chatdev_main.png",
                    "Table 1 主实验：Completeness / Executability / Consistency / Quality 四项全部最高（p.6）"),
                   ("tbl_chatdev_ablation.png",
                    "Table 4 消融：去掉 CDH（去幻觉）与去掉角色分工的对比（p.7）")],
            full=True)),

    dict(
        name="MacNet", accent=TEAL, section="协作机制",
        toc_sub="协作机制 · 规模扩展",
        en_title="Scaling Large Language Model-based\nMulti-Agent Collaboration",
        venue="ICLR 2025 · arXiv:2406.07155",
        authors="Chen Qian, Zihao Xie, YiFei Wang, Wei Liu, Kunlun Zhu 等　（清华大学等）",
        one_liner="把「多智能体能扩到多大」变成可测量的规律：斜 S 形增长、随机拓扑最优。",
        motivation=dict(
            headline="加神经元有效，那加 Agent 呢？",
            purpose="用有向无环图组织 Agent，系统研究「Agent 数量」与「网络拓扑」对任务质量的影响。",
            points=[
                "神经网络的 scaling law 表明增加神经元能持续提升性能，但持续增加 Agent 是否同理未知。",
                "Agent 两两交互时上下文长度随 n² 增长，时间与成本平方级爆炸。",
                "此前研究都在 3–5 个 Agent 的小规模，缺少可用的规模规律。",
            ],
            challenges=[
                "如何在大规模下控制上下文长度，使其不再随规模爆炸。",
                "如何在不预设角色分工的前提下，让 Agent 之间形成有效的批评—改进链。",
                "如何在数量、形状、密度三个维度之间权衡——没有统一最优解。",
                "如何验证性能提升确实来自协作，而不是指标口径造成的假象。",
            ]),
        model=dict(
            headline="DAG 组织 Agent；上下文长度从平方级降到线性级",
            bullets=[
                "MACNET：用有向无环图组织 Agent，节点是 Agent，有向边是一次「批评—改进」推理交互。",
                "交互沿拓扑顺序编排：上游产出初稿，下游逐层反思细化，最后汇聚成完整产物。",
                "不预设角色分工，只做功能二分：actor 负责产出，critic 负责提意见。",
                "六种代表性拓扑：Chain（类瀑布）、Tree、Star、Layer、Mesh、Random。",
                "上下文长度解耦：把增长从平方级降到线性级，这是能扩到千级 Agent 的关键。",
                "没有一个拓扑在所有任务上都最好：tree 在创意写作上明显占优（0.7718）；论文认为 chain 更贴合软件开发的线性流程。",
            ],
            figs=[("macnet_topo.png", "Figure 2/3 六种拓扑；节点放 actor、边放 critic（ICLR 2025, p.3）")]),
        results=dict(
            headline="不规则拓扑优于规则拓扑；协作性能呈斜 S 形增长",
            main=[
                "在多数指标上，MacNet 的各拓扑都超过单 Agent 与已有 MAS 基线（COT / AutoGPT / GPTSwarm / AgentVerse）。",
                "不规则拓扑（Random）最好，Quality 0.6522 为全场最高。",
                "直觉上最密的 Mesh 并不是最优（0.6316）：交互过密会造成信息过载，妨碍反思与细化。",
                "拓扑要按任务选：chain 更适合软件开发（SRDD 0.8056），tree 更适合创意写作（CommonGen 0.7718）。",
            ],
            abl=[
                "节点数从 2⁰ 增到 2⁶：性能先缓升、再快速提升、最后饱和，服从 sigmoid 变体。",
                "节点量级 2⁴ 是性价比合理的选择。",
                "协作涌现比神经网络的涌现更早；critic 提改进后 actor 有 93.10% 概率真的实现。",
            ],
            items=[("tbl_macnet_main.png", "Table 1 主实验：四类基线与六种拓扑的完整对比（p.6）")],
            full=False)),

    dict(
        name="AgentVerse", accent=TEAL, section="协作机制",
        toc_sub="协作机制 · 动态组队",
        en_title="Facilitating Multi-Agent Collaboration\nand Exploring Emergent Behaviors",
        venue="ICLR 2024 · arXiv:2308.10848",
        authors="Weize Chen, Yusheng Su, Jingwei Zuo, Cheng Yang, Chenfei Yuan 等　（清华大学、腾讯微信 AI）",
        one_liner="团队是活的：recruiter 按目标现场招人，评估反馈决定去留。",
        motivation=dict(
            headline="团队怎么组？人工指定角色无法规模化",
            purpose="提出能自动编排专家团队的多智能体框架，并系统观察协作中涌现的行为。",
            points=[
                "多智能体优于单智能体已被反复验证，但「团队怎么组」仍是开放问题。",
                "人工指定角色要求作者事先懂任务，任务一换就得重新设计。",
                "固定团队无法适应任务推进中变化的需求。",
            ],
            challenges=[
                "如何让 Agent 在运行时根据目标生成合适的专家描述，而不是从一个白名单里挑。",
                "如何判断当前团队是否够用，并在不够时增删成员。",
                "如何评估「协作行为本身」，而不只是任务最终得分。",
                "涌现行为不受设计者控制，其中可能包含有害行为。",
            ]),
        model=dict(
            headline="四阶段循环：招募 → 决策 → 执行 → 评估",
            bullets=[
                "Expert Recruitment 专家招募：一个 agent 扮演 recruiter（像 HR），根据目标动态生成一组专家描述，不用预定义角色池。",
                "Collaborative Decision-Making 协同决策：招来的专家一起讨论，产出集体决策。",
                "Action Execution 行动执行：把集体决策落到环境里执行，部分 agent 不一定执行动作。",
                "Evaluation 评估：反馈机制比较当前状态与目标，产出自然语言反馈，驱动下一轮调整（换人或改策略）。",
                "团队规模与成员随任务推进变化——这是「团队是活的」的含义。",
                "三类涌现行为：volunteer（主动帮忙）、conformity（从众收敛）、destructive（破坏性）。",
            ],
            figs=[("agentverse_f1.png", "Figure 1 四阶段循环与多轮团队变化（ICLR 2024, p.2）")]),
        results=dict(
            headline="团队优于单专家与 CoT；同时暴露出破坏性行为风险",
            main=[
                "GPT-4 下 Group 在多数任务上优于 Solo 与 CoT 基线。",
                "逻辑推理（Logic Grid Puzzles）：CoT 59.5 → Solo 64.0 → Group 66.5。",
                "创意写作（Commongen-Challenge）：CoT 95.9 → Solo 99.0。",
                "说明「换成更好的角色描述」本身就有收益，团队协作的增量要按任务看。",
            ],
            abl=[
                "volunteer 与 conformity 发生在决策阶段（对话里）；destructive 绕过决策、直接发生在执行阶段。",
                "destructive 的例子：为拿材料杀掉队友捡掉落物；不去采集而直接拆掉村里的图书馆。",
                "原因是目标优化压过规则遵守；局限：GPT-3.5 下 Group 有时反而不如 Solo。",
            ],
            items=[("tbl_agentverse_main.png",
                    "Table 1 主实验：GPT-3.5 / GPT-4 下 CoT、Solo、Group 三档对比（p.4）")],
            full=False)),

    dict(
        name="Multiagent Debate", accent=TEAL, section="协作机制",
        toc_sub="协作机制 · 对抗式辩论",
        en_title="Improving Factuality and Reasoning in\nLanguage Models through Multiagent Debate",
        venue="ICML 2024 · arXiv:2305.14325",
        authors="Yilun Du, Shuang Li, Antonio Torralba, Joshua B. Tenenbaum, Igor Mordatch　（MIT CSAIL / Google Brain）",
        one_liner="唯一非合作的协作：用分歧做校验，多轮辩论后收敛，降低幻觉。",
        motivation=dict(
            headline="自省有天花板，但多个实例的分歧可以用来纠错",
            purpose="让多个模型实例互相看答案与推理并多轮辩论，收敛到共同答案，以提升推理准确性与事实正确性。",
            points=[
                "单模型会给出错误答案，而且自己检查不出来。",
                "已有做法（CoT、self-consistency、self-reflection）都是让同一个模型换个姿势再想一遍。",
                "同一模型的不同实例答案差异很大，这种分歧尚未被利用。",
            ],
            challenges=[
                "如何让多个实例在有限轮次内收敛到共识，而不是各说各话。",
                "多个 Agent 的完整答案直接拼接会超出上下文长度。",
                "如何确保辩论中传递的是「推理过程」，而不只是结论。",
                "方法必须能直接套在黑盒模型上，不能依赖模型权重或微调。",
            ]),
        model=dict(
            headline="多轮辩论 + 收敛；默认 3 个 agent × 2 轮",
            bullets=[
                "初始轮：n 个 agent（可同一模型、也可不同模型）各自独立回答同一问题。",
                "辩论轮：每个 agent 看到其他 agent 的答案与推理过程，给出修正后的答案。",
                "重复 r 轮后取共识；默认配置为 3 个 agent × 2 轮辩论。",
                "长答案要先摘要：agent 变多时上下文会超限，先用摘要器压缩多方回答再分发。",
                "推理过程必须一起传：只看答案不看推理，效果差很多。",
                "这是唯一的非合作协作：agent 之间互相质疑，而不是分工把事做完。",
            ],
            figs=[],
            side_title="与其它路线的区别",
            side_items=[
                "协作类型：合作（分工） vs 本篇的对抗（互相质疑）",
                "角色：有分工 vs 无分工，同构 agent 做同一件事",
                "传递内容：文档 / 对话 / 决策 vs 答案 + 推理过程",
                "目标：产出更完整的产物 vs 纠错，提高答案正确性",
            ]),
        results=dict(
            headline="六类任务全面超过单模型基线，且能纠正矛盾事实",
            items=[("debate_results.png",
                    "Figure 1 六基准：Single Model（蓝）vs Multi-Model Debate（红）（p.2）"),
                   ("tbl_debate_reasoning.png",
                    "Table 1 推理任务：算术、小学数学、棋步预测的对比（p.6）")],
            full=True)),

    dict(
        name="GeoLLM-Squad", accent=TEAL, section="遥感落地",
        toc_sub="遥感落地 · 多智能体工作流",
        en_title="Multi-Agent Geospatial Copilots\nfor Remote Sensing Workflows",
        venue="arXiv:2501.16254 · UT Austin / SIU / Microsoft",
        authors="Chaehong Lee, Varatheepan Paramanayakam, Andreas Karatzas 等",
        one_liner="把编排与求解分离：在真实遥感工作流上把正确率从 43% 提到 60%。",
        motivation=dict(
            headline="单体 LLM 撑不住遥感工作流：上下文与工具规模都是瓶颈",
            purpose="把「智能体编排」与「地理空间任务求解」拆开，用专职 sub-agent 分摊工具集，突破单体 LLM 的上下文瓶颈。",
            points=[
                "遥感工作流需要多种数据、工具与隐性专业知识，例如云量超阈值时要用 SAR 替代光学产品。",
                "这类「SAR-over-EO」的条件逻辑很难靠提示词硬塞给单体 Copilot。",
                "工具规模从几十涨到几百，单体 Agent 的上下文装不下。",
            ],
            challenges=[
                "如何把 521 个 API 工具分散到多个 Agent，而不让任何单个上下文过载。",
                "如何生成可靠的执行顺序：谁先跑、谁依赖谁的输出。",
                "如何在多轮执行后判断任务是否完成，并决定是否需要重排。",
                "如何让系统在开源小模型上也能工作，以降低云端推理成本。",
            ]),
        model=dict(
            headline="编排与求解分离：一个调度员 + 一圈平级执行者",
            bullets=[
                "核心思路：主 Agent 只负责拆解与调度，具体任务交给专职 sub-agent。",
                "编排者输出「程序式排班表」：指定哪些 Agent、按什么顺序、每个的 sub-prompt。",
                "例：Database(加载 NDVI) → DataOps(筛选 Brisbane) → Agriculture(推荐轮作区) → Map(出图)。",
                "按表执行后，编排者汇总返回消息检查完成度；未完成就改排班表重跑。",
                "Agent 之间零通信：每个 Agent 只跟编排者交换消息——是星形而非多层层级。",
                "521 个 API 工具分散到各专职 Agent，每个只装自己的工具——这是能扛住复杂度增长的原因。",
                "建在 AutoGen（后端）+ GeoLLM-Engine（前端）之上。",
            ],
            figs=[("geollm_arch.png", "Figure 1 各类专职 Agent 与编排器（arXiv:2501.16254, p.2）")]),
        results=dict(
            headline="正确率 43% → 60%；任务变复杂时单 Agent 失效、多 Agent 稳住",
            items=[("geollm_scaling.png",
                    "Figure 2 单 vs 多智能体扩展性消融：任务超过 3 个领域后单 Agent 失效（p.4）"),
                   ("tbl_geollm_main.png",
                    "Table 3 主实验：与单 Agent 及 Chameleon / Magentic 的完整对比（p.3）")],
            full=True)),
]


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = SW, SH

    for i, p in enumerate(PAPERS):
        p["start"] = 4 + 4 * i

    p_cover(prs, 1)
    p_yan(prs, 2)
    p_toc(prs, 3, PAPERS)

    page = 4
    for p in PAPERS:
        p_cover_paper(prs, page, p); page += 1
        p_motivation(prs, page, p); page += 1
        p_model(prs, page, p); page += 1
        if p["results"]["full"]:
            p_results_full(prs, page, p)
        else:
            p_results_split(prs, page, p)
        page += 1

    p_thanks(prs, page)
    prs.save(OUT)
    print("saved:", OUT)
    print("slides:", len(prs.slides._sldIdLst))
    return OUT


if __name__ == "__main__":
    build()
