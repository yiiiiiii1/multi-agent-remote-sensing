# -*- coding: utf-8 -*-
"""交付前自检：扫 PPT 上不该出现的东西。

PPT 上只放「内容」，不放「讲法」和「关系」。这个脚本把所有幻灯片和表格的文字
抽出来，命中违规词就报错退出（非 0），可以挂在构建流程后面。

用法： python3 check_ppt.py [文件名]
"""
import re
import sys
import os
from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT = os.path.join(HERE, "卫星多智能体安全调研汇报.pptx")

# (正则, 说明)
RULES = [
    (r"导师|老师",            "提到老师/导师"),
    (r"★",                    "★ 号"),
    (r"讲法|过渡[:：]|预答|预埋|汇报时|讲的时候", "讲法提示"),
    (r"待讨论|未验证|未经验证", "工作痕迹"),
    (r"v[0-9]\s*修正|v[0-9] 改动|已逐篇核实|全篇核心页", "工作痕迹"),
    (r"注意[:：]|必须区分|别说成|要注明|这里要说明", "提醒/注意事项"),
    (r"我这次汇报|下面每一篇|本篇四页", "关于汇报本身的叙述"),
    (r"谁都不会|不会报警",     "自加的评论"),
    (r"\b[A-E][0-9]\b",         "自编编号（PPT 上不许出现）"),
]


def texts(prs):
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in para.runs).strip()
                    if t:
                        yield i, t
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells:
                        t = cell.text.strip()
                        if t:
                            yield i, t


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT
    prs = Presentation(path)
    hits = 0
    for page, t in texts(prs):
        for pat, why in RULES:
            m = re.search(pat, t)
            if m:
                print(f"  页{page} [{why}] 命中「{m.group(0)}」 → {t[:64]}")
                hits += 1
    print(f"共 {len(prs.slides._sldIdLst)} 页，违规 {hits} 处")
    sys.exit(1 if hits else 0)


if __name__ == "__main__":
    main()
